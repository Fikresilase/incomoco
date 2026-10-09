import io

import docx

from app.domain.models import LLMUsage
from app.rag.generation.answer import AnswerStream
from app.rag.generation.speech_text import SentenceSplitter, split_sentences, to_speakable
from app.rag.ingestion.converters.docx import docx_to_markdown
from app.rag.ingestion.normalizer import normalize_for_search
from app.rag.language import detect_document_lang, detect_lang


def test_language_detection():
    assert detect_lang("ኢንኮሞኮ ምን አገልግሎቶች ይሰጣል?") == "am"
    assert detect_lang("What loans does Inkomoko offer?") == "en"
    assert detect_lang("Inkomoko ብድር ይሰጣል?") == "am"
    assert detect_document_lang("ሰላም ነው። " * 10 + "hello") == "am"
    assert detect_document_lang("hello world " * 5 + "ሰላም " * 5) == "mixed"


def test_amharic_normalization_unifies_homophones_and_punctuation():
    assert normalize_for_search("ሐሳብ") == normalize_for_search("ሀሳብ")
    assert normalize_for_search("ሠላም") == normalize_for_search("ሰላም")
    assert normalize_for_search("ዐይን") == normalize_for_search("አይን")
    assert normalize_for_search("ፀሐይ") == normalize_for_search("ጸሀይ")
    assert normalize_for_search("ሰላም፡ነው።") == "ሰላም ነው"
    assert normalize_for_search("Loans, Training & COACHING!") == "loans training coaching"


def test_to_speakable_strips_markdown_and_citations():
    assert (
        to_speakable("**Loans** are available [1][2]. See [site](http://x).")
        == "Loans are available. See site."
    )


def test_sentence_splitter_emits_complete_sentences_while_streaming():
    splitter = SentenceSplitter(min_chars=10)
    out = []
    for delta in ["Inkomoko offers ", "business loans. ", "It also provides training", " and coaching። Next"]:
        out += splitter.feed(delta)
    out += splitter.flush()
    assert out == ["Inkomoko offers business loans.", "It also provides training and coaching።", "Next"]


def test_split_sentences_groups_for_tts():
    pieces = split_sentences("One sentence here. " * 40, max_chars=200)
    assert all(len(p) <= 200 for p in pieces)
    assert len(pieces) > 1


class _ScriptedLLM:
    model = "scripted"

    def __init__(self, deltas):
        self.deltas = deltas

    async def stream(self, messages, *, temperature=0.2, max_tokens=None, usage: LLMUsage | None = None):
        for d in self.deltas:
            yield d


async def _collect(deltas):
    stream = AnswerStream(_ScriptedLLM(deltas), 0.2)
    text = "".join([d async for d in stream.stream("q", [], [], "en", voice=False)])
    return text, stream.no_answer


async def test_answer_stream_detects_no_answer_marker_split_across_deltas():
    text, no_answer = await _collect(["[[NO_", "ANSWER]] Sorry, ", "I don't know."])
    assert no_answer is True
    assert text == "Sorry, I don't know."


async def test_answer_stream_passes_normal_answers_through():
    text, no_answer = await _collect(["Inkomoko ", "offers loans [1]."])
    assert no_answer is False
    assert text == "Inkomoko offers loans [1]."


async def test_answer_stream_keeps_text_starting_with_bracket():
    text, no_answer = await _collect(["[1] says ", "yes."])
    assert no_answer is False
    assert text == "[1] says yes."


def test_docx_to_markdown_keeps_headings_lists_and_tables_in_order():
    document = docx.Document()
    document.add_heading("Loan Policy", level=1)
    document.add_paragraph("Intro paragraph.")
    document.add_heading("Eligibility", level=2)
    document.add_paragraph("Registered business", style="List Bullet")
    document.add_paragraph("Six months trading", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Product", "Rate"
    table.cell(1, 0).text, table.cell(1, 1).text = "Micro", "12%"
    buffer = io.BytesIO()
    document.save(buffer)

    md = docx_to_markdown(buffer.getvalue())
    assert md.index("# Loan Policy") < md.index("## Eligibility") < md.index("| Product | Rate |")
    assert "- Registered business\n- Six months trading" in md
    assert "| Micro | 12% |" in md


def test_pcm_wrapped_as_wav_and_clips_join_into_one_valid_wav():
    import wave

    from app.core.audio import join_wav, pcm_to_wav, wav_duration_ms

    clip = pcm_to_wav(bytes(48_000))  # 1 s of 24 kHz mono PCM16
    assert clip[:4] == b"RIFF" and wav_duration_ms(clip) == 1000
    joined = join_wav([clip, clip, b""])
    with wave.open(io.BytesIO(joined)) as w:
        assert (w.getframerate(), w.getnchannels(), w.getsampwidth()) == (24_000, 1, 2)
    assert wav_duration_ms(joined) == 2000


def test_router_parsing_defaults_to_search_on_bad_output():
    from app.rag.router import _parse

    assert _parse('{"route": "direct", "say": "Hi there!"}').route == "direct"
    assert _parse('```json\n{"route": "search", "say": "Let me look."}\n```').say == "Let me look."
    assert _parse('{"route": "direct", "say": ""}') is None  # direct needs a reply
    assert _parse('{"route": "maybe"}') is None
    assert _parse("not json") is None


def test_transcription_prompt_has_vocabulary_context_and_language():
    from app.domain.models import ChatTurn
    from app.rag.transcription import VocabularyTerm, build_transcription_prompt

    vocab = [VocabularyTerm.parse("ኢንኮሞኮ=Inkomoko"), VocabularyTerm.parse("ብድር")]
    am = build_transcription_prompt(prompt_lang="am", vocabulary=vocab, context=[ChatTurn("user", "Loans?")])
    assert am.startswith("ይህንን ኦዲዮ ቃል በቃል ይጻፉ።")
    assert "- ኢንኮሞኮ (አማርኛ) / Inkomoko (English)" in am and "- ብድር" in am
    assert "do NOT transcribe" in am and "User: Loans?" in am
    en = build_transcription_prompt(prompt_lang="en", vocabulary=[], context=None)
    assert en.startswith("Transcribe this audio verbatim.") and "Vocabulary" not in en
